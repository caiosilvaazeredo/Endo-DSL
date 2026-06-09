#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o slide deck do exame de qualificação do Endo-DSL (PESC/COPPE/UFRJ).

Uso:
    python presentation/build_deck.py

Produz: presentation/Endo-DSL-Qualificacao.pptx (>= 32 slides), em português,
com tema visual consistente (barra de título slate, acentos índigo), numeração
de slides, rodapé institucional e diagramas simples montados com formas pptx.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn

# --------------------------------------------------------------------------- #
# Paleta e constantes de tema
# --------------------------------------------------------------------------- #
SLATE      = RGBColor(0x1E, 0x29, 0x3B)   # dark slate — barras de título
SLATE_LT   = RGBColor(0x33, 0x41, 0x55)
INDIGO     = RGBColor(0x4F, 0x46, 0xE5)   # acento índigo
INDIGO_LT  = RGBColor(0x81, 0x7C, 0xF0)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
NEARWHITE  = RGBColor(0xF8, 0xFA, 0xFC)
INK        = RGBColor(0x1E, 0x29, 0x3B)
GRAY       = RGBColor(0x64, 0x74, 0x8B)
LIGHTGRAY  = RGBColor(0xE2, 0xE8, 0xF0)
CARD       = RGBColor(0xF1, 0xF5, 0xF9)

# Cores dos níveis de Bloom (conforme README do projeto)
BLOOM = [
    ("Lembrar",     RGBColor(0x25, 0x63, 0xEB)),  # Azul
    ("Compreender", RGBColor(0x0D, 0x94, 0x88)),  # Teal
    ("Aplicar",     RGBColor(0x16, 0xA3, 0x4A)),  # Verde
    ("Analisar",    RGBColor(0xCA, 0x8A, 0x04)),  # Amarelo
    ("Avaliar",     RGBColor(0xEA, 0x58, 0x0C)),  # Laranja
    ("Criar",       RGBColor(0xDC, 0x26, 0x26)),  # Vermelho
]

FONT       = "Calibri"
MONO       = "Consolas"
FOOTER     = "Endo-DSL · PESC/COPPE/UFRJ · Caio Azeredo · 2026"

SW, SH = Inches(13.333), Inches(7.5)   # 16:9 widescreen

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]

# numerador global de slides
_slide_no = {"n": 0}


# --------------------------------------------------------------------------- #
# Utilitários de desenho
# --------------------------------------------------------------------------- #
def _set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _no_fill(shape):
    shape.fill.background()
    shape.line.fill.background()


def box(slide, l, t, w, h, fill=None, line=None, line_w=None, rounded=False, shadow=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        l, t, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w or Pt(1)
    shp.shadow.inherit = False
    if shadow:
        _soft_shadow(shp)
    return shp


def _soft_shadow(shp):
    spPr = shp._element.spPr
    el = spPr.makeelement(qn('a:effectLst'), {})
    sh = el.makeelement(qn('a:outerShdw'), {
        'blurRad': '60000', 'dist': '25000', 'dir': '5400000', 'rotWithShape': '0'})
    clr = sh.makeelement(qn('a:srgbClr'), {'val': '1E293B'})
    alpha = clr.makeelement(qn('a:alpha'), {'val': '24000'})
    clr.append(alpha)
    sh.append(clr)
    el.append(sh)
    spPr.append(el)


def text(slide, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         wrap=True, space_after=4):
    """runs: list of paragraphs; each paragraph is list of (txt, size, color, bold, font, italic)."""
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(2)
    tf.margin_right = Pt(2)
    tf.margin_top = Pt(1)
    tf.margin_bottom = Pt(1)
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        for spec in para:
            txt, size, color, bold, fnt, italic = (list(spec) + [None]*6)[:6]
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size or 18)
            r.font.color.rgb = color or INK
            r.font.bold = bool(bold)
            r.font.name = fnt or FONT
            r.font.italic = bool(italic)
    return tb


def bullets(slide, l, t, w, h, items, size=18, color=INK, gap=8, marker="▸",
            marker_color=INDIGO):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        lvl = 0
        if isinstance(it, tuple):
            it, lvl = it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        p.alignment = PP_ALIGN.LEFT
        if lvl == 0:
            rm = p.add_run(); rm.text = marker + "  "
            rm.font.size = Pt(size); rm.font.bold = True
            rm.font.color.rgb = marker_color; rm.font.name = FONT
            r = p.add_run(); r.text = it
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = FONT
        else:
            p.level = 1
            rm = p.add_run(); rm.text = "–  "
            rm.font.size = Pt(size-2); rm.font.color.rgb = GRAY; rm.font.name = FONT
            r = p.add_run(); r.text = it
            r.font.size = Pt(size-2); r.font.color.rgb = GRAY; r.font.name = FONT
    return tb


def arrow(slide, x1, y1, x2, y2, color=INDIGO, w=Pt(2.25)):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    cn.line.color.rgb = color
    cn.line.width = w
    line = cn.line._get_or_add_ln()
    tail = line.makeelement(qn('a:tailEnd'),
                            {'type': 'triangle', 'w': 'med', 'len': 'med'})
    line.append(tail)
    cn.shadow.inherit = False
    return cn


def chip(slide, l, t, w, h, label, fill, txt_color=WHITE, size=12, bold=True):
    c = box(slide, l, t, w, h, fill=fill, rounded=True)
    tf = c.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
    tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = txt_color; r.font.name = FONT
    return c


# --------------------------------------------------------------------------- #
# Chrome de slide: barra de título, rodapé, número
# --------------------------------------------------------------------------- #
def new_slide(title=None, kicker=None):
    s = prs.slides.add_slide(BLANK)
    _slide_no["n"] += 1
    # fundo branco
    box(s, 0, 0, SW, SH, fill=WHITE)
    if title is not None:
        # barra de título slate
        box(s, 0, 0, SW, Inches(1.15), fill=SLATE)
        # acento índigo
        box(s, 0, Inches(1.15), SW, Inches(0.06), fill=INDIGO)
        # marca lateral
        box(s, 0, 0, Inches(0.16), Inches(1.15), fill=INDIGO)
        runs = []
        if kicker:
            runs.append([(kicker.upper(), 11, INDIGO_LT, True, FONT, False)])
        runs.append([(title, 27, WHITE, True, FONT, False)])
        text(s, Inches(0.55), Inches(0.12), Inches(11.6), Inches(0.95), runs,
             anchor=MSO_ANCHOR.MIDDLE, space_after=0)
        _footer(s)
    return s


def _footer(s):
    box(s, 0, Inches(7.18), SW, Inches(0.32), fill=NEARWHITE)
    text(s, Inches(0.4), Inches(7.18), Inches(10.5), Inches(0.32),
         [[(FOOTER, 9, GRAY, False, FONT, False)]], anchor=MSO_ANCHOR.MIDDLE)
    # número do slide
    text(s, Inches(12.3), Inches(7.18), Inches(0.8), Inches(0.32),
         [[(str(_slide_no["n"]), 10, INDIGO, True, FONT, False)]],
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def body_area():
    """Retângulo útil de conteúdo (l, t, w, h)."""
    return Inches(0.55), Inches(1.45), Inches(12.23), Inches(5.55)


# --------------------------------------------------------------------------- #
# SLIDE 1 — Capa
# --------------------------------------------------------------------------- #
def slide_capa():
    s = prs.slides.add_slide(BLANK)
    _slide_no["n"] += 1
    box(s, 0, 0, SW, SH, fill=SLATE)
    # faixa diagonal de acento
    box(s, 0, Inches(4.55), SW, Inches(0.10), fill=INDIGO)
    box(s, 0, Inches(4.70), SW, Inches(0.04), fill=INDIGO_LT)
    # selo
    chip(s, Inches(0.9), Inches(0.9), Inches(4.6), Inches(0.5),
         "EXAME DE QUALIFICAÇÃO · DOUTORADO", INDIGO, size=13)
    # título
    text(s, Inches(0.9), Inches(1.9), Inches(11.5), Inches(2.2), [
        [("Endo-DSL", 60, WHITE, True, FONT, False)],
        [("Geração semiautomática de jogos educativos HTML5 a partir de uma "
          "DSL estruturada pela Taxonomia de Bloom", 22, INDIGO_LT, False, FONT, False)],
    ], space_after=10)
    # autoria
    text(s, Inches(0.9), Inches(5.0), Inches(11.5), Inches(1.6), [
        [("Doutorando: ", 18, RGBColor(0xCB,0xD5,0xE1), True, FONT, False),
         ("Caio Azeredo", 18, WHITE, False, FONT, False)],
        [("Orientador(a): ", 18, RGBColor(0xCB,0xD5,0xE1), True, FONT, False),
         ("[Prof(a). Orientador(a)]", 18, WHITE, False, FONT, True)],
        [("PESC — Programa de Engenharia de Sistemas e Computação",
          16, RGBColor(0x94,0xA3,0xB8), False, FONT, False)],
        [("COPPE / Universidade Federal do Rio de Janeiro — UFRJ",
          16, RGBColor(0x94,0xA3,0xB8), False, FONT, False)],
    ], space_after=6)
    text(s, Inches(0.9), Inches(6.85), Inches(11.5), Inches(0.4),
         [[("Rio de Janeiro · Junho de 2026", 14, INDIGO_LT, True, FONT, False)]])


# --------------------------------------------------------------------------- #
# SLIDE 2 — Agenda
# --------------------------------------------------------------------------- #
def slide_agenda():
    s = new_slide("Agenda", "Roteiro da apresentação")
    l, t, w, h = body_area()
    cols = [
        ("01", "Contexto & Problema", "Motivação, lacuna de autoria, questões e hipóteses"),
        ("02", "Objetivos", "Objetivo geral e objetivos específicos da tese"),
        ("03", "Fundamentação", "DSLs · Bloom · design endógeno · RAG+LLM"),
        ("04", "Trabalhos relacionados", "Estado da arte e lacuna identificada"),
        ("05", "Proposta Endo-DSL", "Arquitetura de 5 módulos e seus papéis"),
        ("06", "Jornada & Estudo de caso", "Fluxo do usuário e 'Comparando Frações'"),
        ("07", "Metodologia", "Design Science Research e protocolo experimental"),
        ("08", "Resultados & Cronograma", "Prova de conceito, contribuições e plano"),
    ]
    cw = Inches(5.95); ch = Inches(1.18); gx = Inches(0.33); gy = Inches(0.22)
    for i, (num, ti, sub) in enumerate(cols):
        r, c = divmod(i, 2)
        x = l + c * (cw + gx)
        y = t + r * (ch + gy)
        card = box(s, x, y, cw, ch, fill=CARD, rounded=True, shadow=True)
        box(s, x, y, Inches(0.12), ch, fill=INDIGO)
        chip(s, x + Inches(0.28), y + Inches(0.30), Inches(0.85), Inches(0.58),
             num, SLATE, size=20)
        text(s, x + Inches(1.3), y + Inches(0.16), cw - Inches(1.5), ch - Inches(0.3), [
            [(ti, 17, INK, True, FONT, False)],
            [(sub, 12, GRAY, False, FONT, False)],
        ], anchor=MSO_ANCHOR.MIDDLE, space_after=2)


# --------------------------------------------------------------------------- #
# Layout genérico: título + bullets, opcional painel lateral
# --------------------------------------------------------------------------- #
def content_slide(title, kicker, items, size=18, side=None, side_title=None):
    s = new_slide(title, kicker)
    l, t, w, h = body_area()
    bw = w if side is None else Inches(7.4)
    bullets(s, l, t + Inches(0.1), bw, h, items, size=size, gap=12)
    if side is not None:
        sx = l + bw + Inches(0.35)
        sw = w - bw - Inches(0.35)
        panel = box(s, sx, t, sw, Inches(5.0), fill=SLATE, rounded=True, shadow=True)
        box(s, sx, t, sw, Inches(0.6), fill=INDIGO)
        text(s, sx + Inches(0.25), t, sw - Inches(0.5), Inches(0.6),
             [[(side_title or "Em destaque", 14, WHITE, True, FONT, False)]],
             anchor=MSO_ANCHOR.MIDDLE)
        tb = s.shapes.add_textbox(sx + Inches(0.25), t + Inches(0.75),
                                  sw - Inches(0.5), Inches(4.0))
        tf = tb.text_frame; tf.word_wrap = True
        first = True
        for it in side:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(10)
            r = p.add_run(); r.text = "• " + it
            r.font.size = Pt(14); r.font.color.rgb = NEARWHITE; r.font.name = FONT
    return s


# --------------------------------------------------------------------------- #
# SLIDE 3 — Contexto e motivação
# --------------------------------------------------------------------------- #
def slide_contexto():
    content_slide(
        "Contexto e motivação", "Jogos educativos & autoria",
        [
            "Jogos digitais educativos têm eficácia comprovada para engajamento e "
            "aprendizagem ativa, mas seu desenvolvimento é caro e especializado.",
            "Existe um abismo de autoria: educadores dominam o conteúdo, porém não "
            "programam; programadores não dominam a intenção pedagógica.",
            "Ferramentas atuais focam estética e mecânica, raramente garantindo "
            "alinhamento explícito com objetivos cognitivos de aprendizagem.",
            "Resultado frequente: jogos com 'conteúdo justaposto' — perguntas coladas "
            "sobre uma mecânica genérica, sem endogeneidade.",
            "Oportunidade: LLMs e RAG viabilizam geração assistida, desde que ancorada "
            "em um artefato formal verificável.",
        ],
        side=[
            "Custo e tempo de desenvolvimento",
            "Falta de ponte conteúdo↔código",
            "Ausência de garantia pedagógica",
            "Conteúdo justaposto à mecânica",
            "LLM sem âncora formal = alucinação",
        ],
        side_title="Sintomas do problema",
    )


# --------------------------------------------------------------------------- #
# SLIDE 4 — Problema de pesquisa
# --------------------------------------------------------------------------- #
def slide_problema():
    s = new_slide("Problema de pesquisa", "Enunciado central")
    l, t, w, h = body_area()
    panel = box(s, l, t, w, Inches(1.9), fill=SLATE, rounded=True, shadow=True)
    box(s, l, t, Inches(0.16), Inches(1.9), fill=INDIGO)
    text(s, l + Inches(0.5), t + Inches(0.2), w - Inches(1.0), Inches(1.5), [
        [("“Como permitir que educadores especifiquem jogos educativos "
          "pedagogicamente alinhados — e os gerem automaticamente como protótipos "
          "executáveis — sem dominar programação, preservando a "
          "rastreabilidade cognitiva e a endogeneidade?”",
          21, WHITE, True, FONT, True)]],
        anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, l, t + Inches(2.2), w, Inches(3.0), [
        "Desafio 1 — Expressividade vs. simplicidade: uma linguagem acessível, "
        "porém formal e verificável.",
        "Desafio 2 — Garantia pedagógica: o nível de Bloom deve ser um construto "
        "verificável, não um rótulo decorativo.",
        "Desafio 3 — Geração confiável: LLM ancorado por gramática, biblioteca e "
        "validação iterativa (anti-alucinação).",
        "Desafio 4 — Avaliação comparável: medir qualidade de protótipos automáticos "
        "vs. manuais com o mesmo instrumento.",
    ], size=17, gap=12)


# --------------------------------------------------------------------------- #
# SLIDE 5 — Questões de pesquisa
# --------------------------------------------------------------------------- #
def slide_questoes():
    s = new_slide("Questões de pesquisa", "QP1 – QP4")
    l, t, w, h = body_area()
    qs = [
        ("QP1", "Expressividade", "Uma DSL guiada por Bloom consegue expressar jogos "
         "educativos endógenos de forma acessível a não-programadores?"),
        ("QP2", "Verificabilidade", "É possível verificar automaticamente o alinhamento "
         "pedagógico (sintático e semântico) de uma especificação?"),
        ("QP3", "Geração assistida", "Um pipeline RAG+LLM multi-agente, ancorado na DSL, "
         "gera especificações válidas e pedagogicamente coerentes?"),
        ("QP4", "Qualidade comparada", "Protótipos gerados automaticamente alcançam "
         "qualidade pedagógica comparável aos desenvolvidos manualmente?"),
    ]
    cw = Inches(5.95); ch = Inches(2.45); gx = Inches(0.33); gy = Inches(0.3)
    for i, (code, ti, q) in enumerate(qs):
        r, c = divmod(i, 2)
        x = l + c * (cw + gx); y = t + r * (ch + gy)
        box(s, x, y, cw, ch, fill=CARD, rounded=True, shadow=True)
        box(s, x, y, cw, Inches(0.7), fill=SLATE)
        chip(s, x + Inches(0.25), y + Inches(0.16), Inches(1.05), Inches(0.4),
             code, INDIGO, size=15)
        text(s, x + Inches(1.45), y + Inches(0.1), cw - Inches(1.6), Inches(0.55),
             [[(ti, 16, WHITE, True, FONT, False)]], anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + Inches(0.3), y + Inches(0.9), cw - Inches(0.6), ch - Inches(1.0),
             [[(q, 15, INK, False, FONT, False)]])


# --------------------------------------------------------------------------- #
# SLIDE 6 — Hipóteses
# --------------------------------------------------------------------------- #
def slide_hipoteses():
    content_slide(
        "Hipóteses", "Conjecturas verificáveis",
        [
            "H1 — Uma DSL declarativa com Bloom como construto de primeira classe "
            "torna a intenção pedagógica explícita e verificável estaticamente.",
            "H2 — A validação semântica (matriz de afinidade Bloom×mecânica) detecta "
            "incoerências cognitivas antes da compilação.",
            "H3 — Um pipeline multi-agente com refinamento iterativo eleva a taxa de "
            "especificações válidas na primeira tentativa.",
            "H4 — Protótipos automáticos não apresentam diferença significativa de "
            "qualidade pedagógica frente a protótipos manuais (instrumento de 7-D).",
            "H5 — A reparametrização de domínio preserva a coerência pedagógica do "
            "protótipo, evidenciando endogeneidade estrutural.",
        ],
        size=18,
        side=[
            "H1 → QP1 (expressividade)",
            "H2 → QP2 (verificabilidade)",
            "H3 → QP3 (geração)",
            "H4 → QP4 (qualidade)",
            "H5 → endogeneidade",
        ],
        side_title="Mapa hipótese→QP",
    )


# --------------------------------------------------------------------------- #
# SLIDE 7 — Objetivos
# --------------------------------------------------------------------------- #
def slide_objetivos():
    s = new_slide("Objetivos", "Geral e específicos")
    l, t, w, h = body_area()
    # objetivo geral
    box(s, l, t, w, Inches(1.25), fill=SLATE, rounded=True, shadow=True)
    box(s, l, t, Inches(0.16), Inches(1.25), fill=INDIGO)
    text(s, l + Inches(0.45), t + Inches(0.12), w - Inches(0.9), Inches(1.0), [
        [("OBJETIVO GERAL", 12, INDIGO_LT, True, FONT, False)],
        [("Conceber, implementar e avaliar a plataforma Endo-DSL: uma DSL guiada por "
          "Bloom e um pipeline RAG+LLM que geram protótipos de jogos educativos HTML5 "
          "pedagogicamente alinhados e rastreáveis.", 16, WHITE, False, FONT, False)],
    ], anchor=MSO_ANCHOR.MIDDLE, space_after=4)
    objs = [
        "OE1 — Projetar a DSL e sua gramática EBNF com Bloom como construto de 1ª classe.",
        "OE2 — Implementar parser, validação sintática e semântica (matriz de afinidade).",
        "OE3 — Construir a biblioteca curada de componentes pedagógicos (RAG).",
        "OE4 — Desenvolver o pipeline multi-agente recuperar→gerar→validar→refinar.",
        "OE5 — Implementar o compilador DSL→HTML5 autocontido e a reparametrização.",
        "OE6 — Definir e aplicar o instrumento de avaliação pedagógica de 7 dimensões.",
    ]
    bullets(s, l, t + Inches(1.5), w, Inches(3.6), objs, size=16, gap=10)


# --------------------------------------------------------------------------- #
# SLIDE 8 — Fundamentação: DSLs
# --------------------------------------------------------------------------- #
def slide_dsls():
    content_slide(
        "Fundamentação — DSLs", "Linguagens de domínio específico",
        [
            "DSL: linguagem com expressividade restrita a um domínio, ganhando clareza, "
            "verificabilidade e produtividade frente a linguagens de propósito geral.",
            "DSLs externas possuem sintaxe própria e parser dedicado — caso da Endo-DSL "
            "(gramática EBNF + parser recursivo descendente).",
            "Vantagem central: aproximam a notação do vocabulário do especialista do "
            "domínio (aqui, o educador).",
            "Permitem análise estática: erros de coerência são detectados antes da "
            "execução, viabilizando garantias.",
            "Tensão de projeto: expressividade × acessibilidade — resolvida com limites "
            "explícitos de gramática (RF06).",
        ],
        side=[
            "EBNF (ISO/IEC 14977)",
            "Parser recursivo descendente",
            "Análise estática de coerência",
            "Notação próxima do educador",
            "Limites de gramática explícitos",
        ],
        side_title="Na Endo-DSL",
    )


# --------------------------------------------------------------------------- #
# SLIDE 9 — Pirâmide de Bloom
# --------------------------------------------------------------------------- #
def slide_bloom():
    s = new_slide("Fundamentação — Taxonomia de Bloom revisada", "Seis níveis cognitivos")
    l, t, w, h = body_area()
    # pirâmide construída com trapézios empilhados (base larga -> topo)
    levels = list(enumerate(BLOOM))  # 0..5
    cx = l + Inches(3.4)
    top_y = t + Inches(0.2)
    row_h = Inches(0.82)
    max_w = Inches(6.2)
    min_w = Inches(1.9)
    verbs = {
        "Lembrar": "identificar · reconhecer · listar",
        "Compreender": "explicar · classificar · resumir",
        "Aplicar": "usar · resolver · demonstrar",
        "Analisar": "comparar · diferenciar · organizar",
        "Avaliar": "julgar · criticar · justificar",
        "Criar": "construir · planejar · produzir",
    }
    n = len(levels)
    # desenhar do topo (Criar) para baixo? README lista Lembrar..Criar (base..topo)
    # pirâmide: topo = Criar (mais alto cognitivamente)
    ordered = list(reversed(BLOOM))  # Criar no topo
    for i, (name, color) in enumerate(ordered):
        frac = i / (n - 1)
        wbar = Emu(int(min_w + (max_w - min_w) * frac))
        y = top_y + Emu(int(row_h) * i)
        x = Emu(int(cx + (max_w - wbar) / 2))
        bar = s.shapes.add_shape(MSO_SHAPE.TRAPEZOID, x, y, wbar, row_h)
        bar.rotation = 180
        _set_fill(bar, color)
        bar.shadow.inherit = False
        tf = bar.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = name
        r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = WHITE
        r.font.name = FONT
    # legenda de verbos à direita
    lx = cx + max_w + Inches(0.7)
    text(s, lx, top_y, Inches(0)+ (w - (lx - l)), Inches(0.4),
         [[("Verbos típicos por nível", 14, INDIGO, True, FONT, False)]])
    yy = top_y + Inches(0.5)
    for name, color in BLOOM:
        chip(s, lx, yy, Inches(0.22), Inches(0.22), "", color)
        text(s, lx + Inches(0.35), yy - Inches(0.05), Inches(3.6), Inches(0.5), [
            [(name + ": ", 12, INK, True, FONT, False),
             (verbs[name], 12, GRAY, False, FONT, False)],
        ])
        yy += Inches(0.52)
    text(s, l, t + Inches(5.0), w, Inches(0.5),
         [[("Na Endo-DSL, o nível de Bloom é exigível em objetivos, mecânicas e loops — "
            "construto de primeira classe (RF02).", 13, GRAY, False, FONT, True)]])


# --------------------------------------------------------------------------- #
# SLIDE 10 — Game design endógeno
# --------------------------------------------------------------------------- #
def slide_endogeno():
    s = new_slide("Fundamentação — Game design endógeno", "Aprender É jogar")
    l, t, w, h = body_area()
    # dois cartões contrastantes
    cw = Inches(5.95); ch = Inches(2.6)
    # exógeno
    box(s, l, t, cw, ch, fill=RGBColor(0xFE,0xF2,0xF2), rounded=True, shadow=True)
    box(s, l, t, cw, Inches(0.6), fill=RGBColor(0xDC,0x26,0x26))
    text(s, l + Inches(0.3), t, cw - Inches(0.5), Inches(0.6),
         [[("✗  Exógeno (justaposto)", 16, WHITE, True, FONT, False)]],
         anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, l + Inches(0.25), t + Inches(0.8), cw - Inches(0.5), Inches(1.7), [
        "Conteúdo 'colado' sobre mecânica genérica",
        "Recompensa externa ao aprendizado",
        "Acertar a pergunta ≠ jogar bem",
        "Exemplo: quiz sobre uma corrida de carro",
    ], size=14, gap=8, marker="•", marker_color=RGBColor(0xDC,0x26,0x26))
    # endógeno
    x2 = l + cw + Inches(0.33)
    box(s, x2, t, cw, ch, fill=RGBColor(0xF0,0xFD,0xF4), rounded=True, shadow=True)
    box(s, x2, t, cw, Inches(0.6), fill=RGBColor(0x16,0xA3,0x4A))
    text(s, x2 + Inches(0.3), t, cw - Inches(0.5), Inches(0.6),
         [[("✓  Endógeno (integrado)", 16, WHITE, True, FONT, False)]],
         anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, x2 + Inches(0.25), t + Inches(0.8), cw - Inches(0.5), Inches(1.7), [
        "Conteúdo É a mecânica central",
        "Vencer exige exercer a habilidade-alvo",
        "Jogar bem = aprender",
        "Exemplo: comparar frações para avançar",
    ], size=14, gap=8, marker="•", marker_color=RGBColor(0x16,0xA3,0x4A))
    box(s, l, t + Inches(3.0), w, Inches(1.6), fill=CARD, rounded=True)
    text(s, l + Inches(0.4), t + Inches(3.15), w - Inches(0.8), Inches(1.3), [
        [("Princípio de projeto da Endo-DSL", 15, INDIGO, True, FONT, False)],
        [("A endogeneidade é uma dimensão avaliada explicitamente (instrumento de 7-D) "
          "e um critério de validação semântica: a mecânica deve operacionalizar o "
          "objetivo de aprendizagem declarado, não apenas acompanhá-lo.",
          15, INK, False, FONT, False)],
    ], space_after=6)


# --------------------------------------------------------------------------- #
# SLIDE 11 — RAG + LLM multi-agente
# --------------------------------------------------------------------------- #
def slide_rag():
    content_slide(
        "Fundamentação — RAG + LLM multi-agente", "Geração ancorada",
        [
            "LLMs geram texto fluente, mas alucinam estrutura: sem âncora, produzem "
            "DSL sintaticamente inválida ou pedagogicamente incoerente.",
            "RAG (Retrieval-Augmented Generation): recupera componentes curados da "
            "biblioteca e os injeta no contexto de geração.",
            "Arquitetura multi-agente: papéis especializados (recuperar, gerar, "
            "validar, refinar) cooperam em um ciclo fechado.",
            "Refinamento iterativo: erros do validador realimentam o gerador até "
            "convergir para uma especificação válida.",
            "A gramática formal e a matriz de afinidade funcionam como 'guarda-corpo' "
            "verificável do LLM.",
        ],
        side=[
            "Âncora formal (EBNF)",
            "Recuperação curada (RAG)",
            "Agentes especializados",
            "Loop validar→refinar",
            "Anti-alucinação",
        ],
        side_title="Pilares",
    )


# --------------------------------------------------------------------------- #
# SLIDE 12 — Trabalhos relacionados (tabela)
# --------------------------------------------------------------------------- #
def slide_relacionados():
    s = new_slide("Trabalhos relacionados", "Tabela comparativa")
    l, t, w, h = body_area()
    headers = ["Abordagem", "DSL\nformal", "Bloom\n1ª classe", "Geração\nLLM/RAG",
               "Compila\nexecutável", "Avaliação\npedagógica"]
    rows = [
        ["Engines de jogo (Unity/Godot)", "—", "—", "parcial", "✓", "—"],
        ["Autoria visual (Scratch, etc.)", "—", "—", "—", "✓", "parcial"],
        ["Game DSLs (PuzzleScript, Ceptre)", "✓", "—", "—", "✓", "—"],
        ["LLM-to-game (geração direta)", "—", "—", "✓", "parcial", "—"],
        ["Frameworks de design pedagógico", "parcial", "✓", "—", "—", "✓"],
        ["Endo-DSL (esta proposta)", "✓", "✓", "✓", "✓", "✓"],
    ]
    nrow = len(rows) + 1
    col_w = [Inches(3.5), Inches(1.55), Inches(1.7), Inches(1.7), Inches(1.7), Inches(1.95)]
    x0 = l; y0 = t + Inches(0.1)
    rh = Inches(0.72)
    # cabeçalho
    cx = x0
    for j, head in enumerate(headers):
        box(s, cx, y0, col_w[j], rh, fill=SLATE)
        text(s, cx, y0, col_w[j], rh, [[(head, 12, WHITE, True, FONT, False)]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        cx += col_w[j]
    # linhas
    for i, row in enumerate(rows):
        y = y0 + rh + Emu(int(rh) * i)
        last = (i == len(rows) - 1)
        cx = x0
        for j, cell in enumerate(row):
            if last:
                fill = INDIGO
                tcol = WHITE
            else:
                fill = WHITE if i % 2 == 0 else CARD
                tcol = INK
            box(s, cx, y, col_w[j], rh, fill=fill, line=LIGHTGRAY, line_w=Pt(0.5))
            bold = (j == 0) or last
            if cell == "✓" and not last:
                tcol = RGBColor(0x16, 0xA3, 0x4A); bold = True
            text(s, cx, y, col_w[j], rh, [[(cell, 12.5, tcol, bold, FONT, False)]],
                 align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
            cx += col_w[j]
    text(s, l, y0 + rh * nrow + Inches(0.1), w, Inches(0.5),
         [[("Nenhuma abordagem reúne, simultaneamente, DSL formal, Bloom como construto "
            "de 1ª classe, geração RAG+LLM, compilação executável e avaliação "
            "pedagógica comparável.", 12, GRAY, False, FONT, True)]])


# --------------------------------------------------------------------------- #
# SLIDE 13 — Lacuna
# --------------------------------------------------------------------------- #
def slide_lacuna():
    s = new_slide("Lacuna identificada", "Posicionamento da contribuição")
    l, t, w, h = body_area()
    box(s, l, t, w, Inches(1.4), fill=INDIGO, rounded=True, shadow=True)
    text(s, l + Inches(0.5), t + Inches(0.15), w - Inches(1.0), Inches(1.1),
         [[("Não há uma plataforma que una uma DSL pedagógica verificável a um "
            "pipeline de geração ancorado e a uma avaliação comparável — fechando o "
            "ciclo intenção → especificação → protótipo → evidência.",
            19, WHITE, True, FONT, False)]], anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, l, t + Inches(1.7), w, Inches(3.2), [
        "Game DSLs existem, mas ignoram a dimensão cognitiva (Bloom).",
        "Geradores LLM existem, mas sem âncora formal nem garantia pedagógica.",
        "Instrumentos pedagógicos existem, mas desacoplados da geração de artefatos.",
        "A Endo-DSL integra os três eixos num único artefato rastreável e avaliável.",
    ], size=18, gap=14)


# --------------------------------------------------------------------------- #
# SLIDE 14 — Visão geral da proposta
# --------------------------------------------------------------------------- #
def slide_visao():
    s = new_slide("Visão geral da proposta Endo-DSL", "Do contexto ao protótipo avaliável")
    l, t, w, h = body_area()
    stages = [
        ("Contexto\neducacional", SLATE),
        ("DSL guiada\npor Bloom", INDIGO),
        ("Validação\nsint.+sem.", RGBColor(0x0D,0x94,0x88)),
        ("Compilação\nHTML5", RGBColor(0x16,0xA3,0x4A)),
        ("Avaliação\n7 dimensões", RGBColor(0xEA,0x58,0x0C)),
    ]
    n = len(stages)
    bw = Inches(2.05); bh = Inches(1.4); gap = Inches(0.42)
    total = bw * n + gap * (n - 1)
    x = l + Emu(int((w - total) / 2)); y = t + Inches(0.4)
    cy = Emu(int(y + bh / 2))
    for i, (lab, col) in enumerate(stages):
        bx = x + Emu(int((bw + gap)) * i)
        chip(s, bx, y, bw, bh, lab, col, size=15)
        if i < n - 1:
            arrow(s, bx + bw, cy, bx + bw + gap, cy)
    box(s, l, t + Inches(2.4), w, Inches(2.4), fill=CARD, rounded=True)
    bullets(s, l + Inches(0.4), t + Inches(2.6), w - Inches(0.8), Inches(2.1), [
        "Tudo flui de um único artefato formal: a especificação Endo-DSL.",
        "RAG+LLM podem propor a especificação; o educador a revisa em editor ao vivo.",
        "Rastreabilidade pedagógica preservada do objetivo até o protótipo executável.",
        "Protótipos automáticos e manuais avaliados pelo MESMO instrumento (comparação).",
    ], size=16, gap=11)


# --------------------------------------------------------------------------- #
# SLIDE 15 — Arquitetura (5 módulos)
# --------------------------------------------------------------------------- #
def slide_arquitetura():
    s = new_slide("Arquitetura", "Cinco módulos integrados pela fachada Platform")
    l, t, w, h = body_area()
    mods = [
        ("M1 · Linguagem/DSL", "EBNF · parser RD · validação", INDIGO),
        ("M2 · Biblioteca", "componentes curados · SQLite · RAG", RGBColor(0x0D,0x94,0x88)),
        ("M3 · Pipeline multi-agente", "recuperar→gerar→validar→refinar", RGBColor(0x7C,0x3A,0xED)),
        ("M4 · Compilador HTML5", "jogo autocontido · reparametrização", RGBColor(0x16,0xA3,0x4A)),
        ("M5 · Avaliação", "instrumento Likert 7-D · relatórios", RGBColor(0xEA,0x58,0x0C)),
    ]
    # fachada no topo
    box(s, l, t, w, Inches(0.8), fill=SLATE, rounded=True, shadow=True)
    text(s, l, t, w, Inches(0.8),
         [[("Platform — fachada única (CLI · Web tipo-Overleaf · API HTTP)",
            17, WHITE, True, FONT, False)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # cinco módulos em linha
    n = len(mods); gap = Inches(0.25)
    bw = Emu(int((w - gap * (n - 1)) / n)); bh = Inches(2.1)
    y = t + Inches(1.25)
    for i, (ti, sub, col) in enumerate(mods):
        x = l + Emu(int(bw) * i) + gap * i
        arrow(s, Emu(int(x + bw / 2)), t + Inches(0.8),
              Emu(int(x + bw / 2)), y, color=GRAY, w=Pt(1.5))
        box(s, x, y, bw, bh, fill=WHITE, line=col, line_w=Pt(2.25), rounded=True, shadow=True)
        box(s, x, y, bw, Inches(0.12), fill=col)
        text(s, x + Inches(0.15), y + Inches(0.2), bw - Inches(0.3), bh - Inches(0.4), [
            [(ti, 13.5, col, True, FONT, False)],
            [("", 4, INK, False, FONT, False)],
            [(sub, 11.5, GRAY, False, FONT, False)],
        ], anchor=MSO_ANCHOR.TOP, space_after=3)
    # camada de dados
    yd = y + bh + Inches(0.3)
    box(s, l, yd, w, Inches(0.7), fill=CARD, line=LIGHTGRAY, line_w=Pt(1), rounded=True)
    text(s, l, yd, w, Inches(0.7),
         [[("Persistência: SQLite (sessões · componentes · tentativas · avaliações)  ·  "
            "Workspace de protótipos HTML5", 13, INK, True, FONT, False)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# --------------------------------------------------------------------------- #
# SLIDE 16 — Módulo 1: DSL + EBNF
# --------------------------------------------------------------------------- #
def _code_panel(s, l, t, w, h, lines, title="EBNF"):
    box(s, l, t, w, h, fill=RGBColor(0x0F,0x17,0x2A), rounded=True, shadow=True)
    box(s, l, t, w, Inches(0.42), fill=SLATE_LT)
    text(s, l + Inches(0.25), t, w - Inches(0.5), Inches(0.42),
         [[(title, 12, INDIGO_LT, True, MONO, False)]], anchor=MSO_ANCHOR.MIDDLE)
    tb = s.shapes.add_textbox(l + Inches(0.25), t + Inches(0.55), w - Inches(0.5), h - Inches(0.7))
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    for ln, col in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(2)
        r = p.add_run(); r.text = ln if ln else " "
        r.font.size = Pt(11.5); r.font.name = MONO
        r.font.color.rgb = col
    return tb


def slide_mod1():
    s = new_slide("Módulo 1 — Linguagem / DSL", "Gramática EBNF & exemplo (RF01)")
    l, t, w, h = body_area()
    cw = Inches(5.95)
    green = RGBColor(0x6E,0xE7,0xB7); blue = RGBColor(0x93,0xC5,0xFD)
    gray = RGBColor(0x94,0xA3,0xB8); white = RGBColor(0xE2,0xE8,0xF0)
    ebnf = [
        ("spec      = game_decl ;", white),
        ("game_decl = \"game\" string \"{\"", white),
        ("            { game_member } \"}\" ;", white),
        ("", gray),
        ("mechanic_decl =", white),
        ("   \"mechanic\" ident \"{\"", white),
        ("   { mechanic_field } \"}\" ;", white),
        ("", gray),
        ("(* Bloom = construto 1ª classe *)", gray),
        ("bloom = \"Lembrar\" | \"Compreender\"", green),
        ("      | \"Aplicar\"  | \"Analisar\"", green),
        ("      | \"Avaliar\"  | \"Criar\" ;", green),
    ]
    _code_panel(s, l, t, cw, Inches(5.0), ebnf, title="grammar.ebnf (ISO/IEC 14977)")
    x2 = l + cw + Inches(0.33)
    ex = [
        ("game \"Comparando Frações\" {", blue),
        ("  metadata {", white),
        ("    domain: \"Matemática\"", white),
        ("    bloom: Analisar", green),
        ("    age_range: \"10-11\"", white),
        ("  }", white),
        ("  mechanic comparacao {", blue),
        ("    type: classification", white),
        ("    bloom: Analisar", green),
        ("    addresses: obj_comparar", white),
        ("  }", white),
        ("}", blue),
    ]
    _code_panel(s, x2, t, cw, Inches(5.0), ex, title="fracoes.endo")


# --------------------------------------------------------------------------- #
# SLIDE 17 — Bloom como construto de 1ª classe (RF02)
# --------------------------------------------------------------------------- #
def slide_bloom_firstclass():
    content_slide(
        "Bloom como construto de primeira classe", "RF02",
        [
            "O não-terminal <bloom> é exigível em objetivos, mecânicas, loops e "
            "ramificações narrativas — não é metadado opcional.",
            "Nomes canônicos em português; o parser também aceita os equivalentes em "
            "inglês (Remember…Create).",
            "Cada elemento declara o nível cognitivo que pretende exercitar, tornando a "
            "intenção pedagógica explícita e auditável.",
            "Isso habilita a verificação semântica: comparar o Bloom do objetivo com o "
            "Bloom das mecânicas que o endereçam.",
            "Resultado: rastreabilidade pedagógica do objetivo de aprendizagem até a "
            "mecânica e o protótipo compilado.",
        ],
        size=18,
        side=[lvl[0] for lvl in BLOOM],
        side_title="Os 6 níveis aceitos",
    )


# --------------------------------------------------------------------------- #
# SLIDE 18 — Validação + matriz de afinidade (RF03/RF04)
# --------------------------------------------------------------------------- #
def slide_validacao():
    s = new_slide("Validação sintática e semântica", "RF03 / RF04 — matriz de afinidade")
    l, t, w, h = body_area()
    # esquerda: dois tipos de validação
    cw = Inches(5.2)
    box(s, l, t, cw, Inches(2.2), fill=CARD, rounded=True, shadow=True)
    text(s, l + Inches(0.3), t + Inches(0.15), cw - Inches(0.6), Inches(2.0), [
        [("RF03 · Validação sintática", 16, INDIGO, True, FONT, False)],
        [("Parser recursivo descendente verifica a aderência à gramática EBNF; "
          "erros reportam linha/coluna e sugestão.", 13.5, INK, False, FONT, False)],
        [("RF04 · Validação semântica", 16, INDIGO, True, FONT, False)],
        [("Coerência cognitiva: o nível de Bloom da mecânica é compatível com o do "
          "objetivo que ela endereça (matriz de afinidade).", 13.5, INK, False, FONT, False)],
    ], space_after=8)
    bullets(s, l, t + Inches(2.45), cw, Inches(2.5), [
        "Referências cruzadas (addresses) resolvidas.",
        "Loops com transições válidas entre mecânicas.",
        "Avisos pedagógicos não bloqueiam, mas orientam.",
    ], size=14, gap=9)
    # direita: matriz de afinidade Bloom x Bloom (objetivo vs mecânica)
    mx = l + cw + Inches(0.45)
    names = [b[0][:4] for b in BLOOM]
    cell = Inches(0.72); top = t + Inches(0.55); left = mx + Inches(0.9)
    text(s, mx, t, Inches(6.5), Inches(0.45),
         [[("Matriz de afinidade  (linha: objetivo · coluna: mecânica)",
            12.5, INK, True, FONT, False)]])
    # cabeçalho colunas
    for j, nm in enumerate(names):
        box(s, left + Emu(int(cell) * j), top, cell, Inches(0.4),
            fill=BLOOM[j][1])
        text(s, left + Emu(int(cell) * j), top, cell, Inches(0.4),
             [[(nm, 10, WHITE, True, FONT, False)]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for i in range(6):
        ry = top + Inches(0.4) + Emu(int(cell) * i)
        box(s, mx, ry, Inches(0.9), cell, fill=BLOOM[i][1])
        text(s, mx, ry, Inches(0.9), cell,
             [[(names[i], 10, WHITE, True, FONT, False)]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        for j in range(6):
            cx = left + Emu(int(cell) * j)
            diff = abs(i - j)
            if diff == 0:
                fill = RGBColor(0x16, 0xA3, 0x4A); mark = "✓"
            elif diff == 1:
                fill = RGBColor(0xCA, 0x8A, 0x04); mark = "~"
            else:
                fill = RGBColor(0xDC, 0x26, 0x26); mark = "✗"
            box(s, cx, ry, cell, cell, fill=fill, line=WHITE, line_w=Pt(1))
            text(s, cx, ry, cell, cell, [[(mark, 13, WHITE, True, FONT, False)]],
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    ly = top + Inches(0.4) + Emu(int(cell) * 6) + Inches(0.15)
    text(s, mx, ly, Inches(6.5), Inches(0.5),
         [[("✓ coerente   ~ aviso (1 nível)   ✗ incoerente (≥2 níveis)",
            12, GRAY, False, FONT, False)]])


# --------------------------------------------------------------------------- #
# SLIDE 19 — Módulo 2: biblioteca
# --------------------------------------------------------------------------- #
def slide_mod2():
    content_slide(
        "Módulo 2 — Biblioteca de componentes", "SQLite & curadoria (RF07–RF12)",
        [
            "Repositório de componentes pedagógicos reutilizáveis: mecânicas "
            "parametrizáveis indexadas por nível de Bloom, tipo e domínio.",
            "Persistência em SQLite (stdlib) — zero dependências externas obrigatórias.",
            "Status duplo: componentes canônicos (★, curados) e experimentais (○).",
            "Métricas por componente: instanciações e avaliação média realimentam a "
            "recuperação RAG.",
            "Jornada do curador: aprovar/rejeitar componentes experimentais com "
            "justificativa (RF12), promovendo a canônicos.",
        ],
        size=18,
        side=[
            "Indexação por Bloom/tipo/domínio",
            "SQLite (stdlib)",
            "★ canônico · ○ experimental",
            "Métricas de uso e qualidade",
            "Fila de curadoria",
        ],
        side_title="Características",
    )


# --------------------------------------------------------------------------- #
# SLIDE 20 — Módulo 3: pipeline multi-agente
# --------------------------------------------------------------------------- #
def slide_mod3():
    s = new_slide("Módulo 3 — Pipeline multi-agente", "recuperar→gerar→validar→refinar (RF13–RF18)")
    l, t, w, h = body_area()
    stages = [
        ("Recuperar", "RAG: componentes\nrelevantes (RF14)", RGBColor(0x0D,0x94,0x88)),
        ("Gerar", "LLM produz DSL\nancorada (RF15)", RGBColor(0x4F,0x46,0xE5)),
        ("Validar", "sintaxe + semântica\nBloom (RF16)", RGBColor(0xCA,0x8A,0x04)),
        ("Refinar", "realimenta erros\n(RF17/RF18)", RGBColor(0xEA,0x58,0x0C)),
    ]
    n = len(stages); bw = Inches(2.6); bh = Inches(1.7); gap = Inches(0.55)
    total = bw * n + gap * (n - 1)
    x = l + Emu(int((w - total) / 2)); y = t + Inches(0.35)
    cy = Emu(int(y + bh / 2))
    coords = []
    for i, (ti, sub, col) in enumerate(stages):
        bx = x + Emu(int((bw + gap)) * i)
        coords.append((bx, col))
        box(s, bx, y, bw, bh, fill=col, rounded=True, shadow=True)
        text(s, bx + Inches(0.15), y + Inches(0.18), bw - Inches(0.3), bh - Inches(0.4), [
            [(f"{i+1}. {ti}", 17, WHITE, True, FONT, False)],
            [(sub, 12, RGBColor(0xE2,0xE8,0xF0), False, FONT, False)],
        ], anchor=MSO_ANCHOR.MIDDLE, space_after=4)
        if i < n - 1:
            arrow(s, bx + bw, cy, bx + bw + gap, cy, color=SLATE, w=Pt(2.5))
    # loop de refinamento de volta para "Gerar"
    bx_ref = coords[3][0]; bx_gen = coords[1][0]
    yb = y + bh + Inches(0.55)
    arrow(s, Emu(int(bx_ref + bw/2)), y + bh, Emu(int(bx_ref + bw/2)), yb,
          color=RGBColor(0xDC,0x26,0x26), w=Pt(2))
    cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                Emu(int(bx_ref + bw/2)), yb,
                                Emu(int(bx_gen + bw/2)), yb)
    cn.line.color.rgb = RGBColor(0xDC,0x26,0x26); cn.line.width = Pt(2)
    cn.shadow.inherit = False
    arrow(s, Emu(int(bx_gen + bw/2)), yb, Emu(int(bx_gen + bw/2)), y + bh,
          color=RGBColor(0xDC,0x26,0x26), w=Pt(2))
    text(s, bx_gen, yb + Inches(0.05), bx_ref - bx_gen + bw, Inches(0.4),
         [[("ciclo de refinamento iterativo (até validar ou esgotar tentativas)",
            12, RGBColor(0xDC,0x26,0x26), True, FONT, True)]],
         align=PP_ALIGN.CENTER)
    box(s, l, t + Inches(3.5), w, Inches(1.3), fill=CARD, rounded=True)
    text(s, l + Inches(0.4), t + Inches(3.6), w - Inches(0.8), Inches(1.1), [
        [("Métricas instrumentadas: ", 14, INDIGO, True, FONT, False),
         ("taxa de validade na 1ª tentativa, número médio de tentativas, taxa de "
          "rejeição e erros mais frequentes — registradas em generation_attempts.",
          14, INK, False, FONT, False)],
    ], anchor=MSO_ANCHOR.MIDDLE)


# --------------------------------------------------------------------------- #
# SLIDE 21 — Módulo 4: compilador HTML5
# --------------------------------------------------------------------------- #
def slide_mod4():
    content_slide(
        "Módulo 4 — Compilador HTML5", "RF19–RF23 & reparametrização (RF22)",
        [
            "Compila a especificação DSL em um jogo HTML5 autocontido — zero "
            "dependências externas em runtime (HTML+CSS+JS embutidos).",
            "Cada mecânica da DSL mapeia para um gerador de mecânica de jogo "
            "(quiz, classificação, sequenciamento, comparação…).",
            "Gera artefato de rastreabilidade pedagógica: do objetivo ao componente "
            "executável (/prototype/<pid>/traceability).",
            "Reparametrização de domínio (RF22): troca o domínio/tópico de um protótipo "
            "existente SEM recompilar — evidência de endogeneidade estrutural.",
            "Saída executável em qualquer navegador, distribuível como arquivo único.",
        ],
        size=18,
        side=[
            "HTML5 autocontido (offline)",
            "Mecânica DSL → mecânica jogo",
            "Rastreabilidade pedagógica",
            "Reparametrização sem recompilar",
            "Distribuível como 1 arquivo",
        ],
        side_title="Capacidades",
    )


# --------------------------------------------------------------------------- #
# SLIDE 22 — Módulo 5: avaliação 7-D
# --------------------------------------------------------------------------- #
def slide_mod5():
    s = new_slide("Módulo 5 — Avaliação pedagógica", "Instrumento Likert de 7 dimensões (RF24–RF26)")
    l, t, w, h = body_area()
    dims = [
        ("Alinhamento pedagógico", "Mecânicas implementam os objetivos?", "1,3"),
        ("Coerência cognitiva (Bloom)", "Nível exercitado = nível pretendido?", "1,2"),
        ("Endogeneidade", "Aprender é jogar (não justaposto)?", "1,3"),
        ("Clareza instrucional", "Instruções e desafios claros?", "1,0"),
        ("Adequação ao público", "Adequado à faixa etária/contexto?", "1,0"),
        ("Potencial de engajamento", "Mantém o interesse do aprendiz?", "1,0"),
        ("Adaptabilidade de conteúdo", "Reparametrizável p/ outros domínios?", "0,8"),
    ]
    cw = Inches(5.95); rh = Inches(0.66)
    for i, (ti, q, wgt) in enumerate(dims):
        col = i % 2
        r = i // 2 if col == 0 else (i // 2)
        # simple two-column layout
        x = l + col * (cw + Inches(0.33))
        y = t + (i // 2) * (rh + Inches(0.14))
        box(s, x, y, cw, rh, fill=CARD, rounded=True)
        box(s, x, y, Inches(0.1), rh, fill=INDIGO)
        chip(s, x + cw - Inches(0.85), y + Inches(0.13), Inches(0.7), Inches(0.4),
             "w " + wgt, SLATE, size=11)
        text(s, x + Inches(0.28), y + Inches(0.04), cw - Inches(1.2), rh - Inches(0.08), [
            [(f"D{i+1}. {ti}", 14, INK, True, FONT, False)],
            [(q, 11, GRAY, False, FONT, False)],
        ], anchor=MSO_ANCHOR.MIDDLE, space_after=1)
    yb = t + 4 * (rh + Inches(0.14)) + Inches(0.05)
    box(s, l, yb, w, Inches(0.75), fill=SLATE, rounded=True)
    text(s, l + Inches(0.4), yb, w - Inches(0.8), Inches(0.75),
         [[("Escala Likert 1–5 por dimensão · pesos dialogam com MEEGA+ · "
            "MESMO instrumento para protótipos automáticos e manuais (RF26 → QP4).",
            13, WHITE, True, FONT, False)]], anchor=MSO_ANCHOR.MIDDLE)


# --------------------------------------------------------------------------- #
# SLIDE 23 — Jornada do usuário (stepper)
# --------------------------------------------------------------------------- #
def slide_jornada():
    s = new_slide("Jornada do usuário", "Fluxo do Studio — Fases 1 a 7")
    l, t, w, h = body_area()
    phases = [
        ("1", "Contexto", "domínio, Bloom,\nfaixa, duração"),
        ("2", "Recuperação", "componentes\nvia RAG"),
        ("3", "Geração", "DSL a partir do\ncontexto"),
        ("4", "Revisão", "edição com\nvalidação ao vivo"),
        ("5", "Compilação", "protótipo HTML5\n+ métricas"),
        ("6", "Avaliação", "instrumento\n7-D"),
        ("7", "Curadoria", "promoção de\ncomponentes"),
    ]
    n = len(phases)
    bw = Inches(1.55); bh = Inches(1.55)
    gap = Emu(int((w - bw * n) / (n - 1)))
    y = t + Inches(0.7)
    cy = Emu(int(y + Inches(0.45)))
    for i, (num, ti, sub) in enumerate(phases):
        x = l + Emu(int(bw) * i) + Emu(int(gap)) * i
        col = BLOOM[i % len(BLOOM)][1] if i < 6 else SLATE
        col = INDIGO if i < 6 else RGBColor(0x7C,0x3A,0xED)
        # círculo numerado
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Emu(int((bw-Inches(0.9))/2)), y,
                                  Inches(0.9), Inches(0.9))
        _set_fill(circ, col); circ.shadow.inherit = False
        tf = circ.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        rr = p.add_run(); rr.text = num
        rr.font.size = Pt(26); rr.font.bold = True; rr.font.color.rgb = WHITE
        rr.font.name = FONT
        text(s, x - Inches(0.15), y + Inches(1.0), bw + Inches(0.3), Inches(1.2), [
            [(ti, 14, INK, True, FONT, False)],
            [(sub, 10.5, GRAY, False, FONT, False)],
        ], align=PP_ALIGN.CENTER, space_after=2)
        if i < n - 1:
            nx = l + Emu(int(bw) * (i+1)) + Emu(int(gap)) * (i+1)
            arrow(s, x + Emu(int((bw+Inches(0.9))/2)), cy,
                  nx + Emu(int((bw-Inches(0.9))/2)), cy, color=INDIGO_LT, w=Pt(2))
    box(s, l, t + Inches(3.6), w, Inches(1.1), fill=CARD, rounded=True)
    text(s, l + Inches(0.4), t + Inches(3.7), w - Inches(0.8), Inches(0.9),
         [[("As Fases 1–6 percorrem o Studio (interface tipo-Overleaf); a Fase 7 "
            "corresponde à jornada do curador, que governa a evolução da biblioteca.",
            14, INK, False, FONT, False)]], anchor=MSO_ANCHOR.MIDDLE)


# --------------------------------------------------------------------------- #
# SLIDE 24 — Interface tipo-Overleaf
# --------------------------------------------------------------------------- #
def slide_interface():
    s = new_slide("Interface tipo-Overleaf", "Editor dividido + live preview")
    l, t, w, h = body_area()
    # mockup: editor à esquerda, preview à direita
    cw = Emu(int((w - Inches(0.4)) / 2))
    # editor
    box(s, l, t, cw, Inches(4.3), fill=RGBColor(0x0F,0x17,0x2A), rounded=True, shadow=True)
    box(s, l, t, cw, Inches(0.45), fill=SLATE_LT)
    text(s, l + Inches(0.25), t, cw - Inches(0.5), Inches(0.45),
         [[("◳  Editor DSL", 12, INDIGO_LT, True, FONT, False)]],
         anchor=MSO_ANCHOR.MIDDLE)
    edlines = [
        ('game "Comparando Frações" {', RGBColor(0x93,0xC5,0xFD)),
        ('  metadata { bloom: Analisar }', RGBColor(0x6E,0xE7,0xB7)),
        ('  mechanic comparacao {', RGBColor(0xE2,0xE8,0xF0)),
        ('    type: classification', RGBColor(0xE2,0xE8,0xF0)),
        ('    bloom: Analisar }', RGBColor(0x6E,0xE7,0xB7)),
        ('}', RGBColor(0x93,0xC5,0xFD)),
    ]
    tb = s.shapes.add_textbox(l + Inches(0.25), t + Inches(0.6), cw - Inches(0.5), Inches(3.5))
    tf = tb.text_frame; tf.word_wrap = True; first = True
    for ln, col in edlines:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.space_after = Pt(4)
        r = p.add_run(); r.text = ln; r.font.size = Pt(13); r.font.name = MONO
        r.font.color.rgb = col
    # preview
    x2 = l + cw + Inches(0.4)
    box(s, x2, t, cw, Inches(4.3), fill=WHITE, line=LIGHTGRAY, line_w=Pt(1.5),
        rounded=True, shadow=True)
    box(s, x2, t, cw, Inches(0.45), fill=INDIGO)
    text(s, x2 + Inches(0.25), t, cw - Inches(0.5), Inches(0.45),
         [[("▶  Live preview (protótipo HTML5)", 12, WHITE, True, FONT, False)]],
         anchor=MSO_ANCHOR.MIDDLE)
    # mock do jogo
    box(s, x2 + Inches(0.4), t + Inches(0.75), cw - Inches(0.8), Inches(0.7),
        fill=CARD, rounded=True)
    text(s, x2 + Inches(0.4), t + Inches(0.75), cw - Inches(0.8), Inches(0.7),
         [[("Qual é maior?   3/4  vs  2/3", 15, INK, True, FONT, False)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for i, (lab, col) in enumerate([("3/4 maior", RGBColor(0x16,0xA3,0x4A)),
                                    ("2/3 maior", RGBColor(0xDC,0x26,0x26)),
                                    ("iguais", GRAY)]):
        chip(s, x2 + Inches(0.4) + Emu(int(Inches(1.65)) * i), t + Inches(1.65),
             Inches(1.5), Inches(0.5), lab, col, size=12)
    bullets(s, x2 + Inches(0.4), t + Inches(2.4), cw - Inches(0.8), Inches(1.7), [
        "Validação ao vivo (linha/coluna).",
        "Recompila a cada alteração.",
        "Sem instalar nada: roda no navegador.",
    ], size=12.5, gap=7)
    bullets(s, l, t + Inches(4.5), w, Inches(0.8), [
        "Metáfora familiar (Overleaf/LaTeX): código formal à esquerda, artefato "
        "renderizado à direita — baixa curva para o educador.",
    ], size=14, gap=6)


# --------------------------------------------------------------------------- #
# SLIDE 25 — Estudo de caso
# --------------------------------------------------------------------------- #
def slide_estudo():
    s = new_slide("Estudo de caso — “Comparando Frações”", "Bloom: Analisar")
    l, t, w, h = body_area()
    # contexto chips
    ctx = [("Domínio", "Matemática"), ("Tópico", "Frações"),
           ("Bloom", "Analisar"), ("Faixa", "10–11 anos"), ("Duração", "15 min")]
    cwid = Inches(2.3)
    for i, (k, v) in enumerate(ctx):
        x = l + Emu(int(cwid + Inches(0.1)) * i)
        box(s, x, t, cwid, Inches(0.9), fill=CARD, rounded=True)
        box(s, x, t, cwid, Inches(0.3), fill=SLATE)
        text(s, x, t, cwid, Inches(0.3), [[(k, 11, WHITE, True, FONT, False)]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x, t + Inches(0.32), cwid, Inches(0.55),
             [[(v, 14, INK, True, FONT, False)]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # fluxo objetivo->mecanicas->loop
    y2 = t + Inches(1.3)
    box(s, l, y2, Inches(3.9), Inches(1.1), fill=RGBColor(0xCA,0x8A,0x04), rounded=True, shadow=True)
    text(s, l + Inches(0.2), y2, Inches(3.5), Inches(1.1), [
        [("OBJETIVO obj_comparar", 12, WHITE, True, FONT, False)],
        [("Comparar frações com denom.\ndiferentes (Analisar)", 12, WHITE, False, FONT, False)],
    ], anchor=MSO_ANCHOR.MIDDLE, space_after=2)
    arrow(s, l + Inches(3.9), Emu(int(y2 + Inches(0.55))),
          l + Inches(4.5), Emu(int(y2 + Inches(0.55))))
    box(s, l + Inches(4.5), y2, Inches(3.9), Inches(1.1), fill=INDIGO, rounded=True, shadow=True)
    text(s, l + Inches(4.7), y2, Inches(3.5), Inches(1.1), [
        [("MECÂNICA comparacao", 12, WHITE, True, FONT, False)],
        [("classification · Analisar\nclassifica pares: maior/menor/igual", 11.5, WHITE, False, FONT, False)],
    ], anchor=MSO_ANCHOR.MIDDLE, space_after=2)
    arrow(s, l + Inches(8.4), Emu(int(y2 + Inches(0.55))),
          l + Inches(9.0), Emu(int(y2 + Inches(0.55))))
    box(s, l + Inches(9.0), y2, Inches(3.2), Inches(1.1),
        fill=RGBColor(0x25,0x63,0xEB), rounded=True, shadow=True)
    text(s, l + Inches(9.2), y2, Inches(2.9), Inches(1.1), [
        [("MECÂNICA revisao", 12, WHITE, True, FONT, False)],
        [("quiz · Lembrar\nfixação dos conceitos", 11.5, WHITE, False, FONT, False)],
    ], anchor=MSO_ANCHOR.MIDDLE, space_after=2)
    # loop
    y3 = y2 + Inches(1.4)
    box(s, l, y3, w, Inches(0.85), fill=RGBColor(0x0D,0x94,0x88), rounded=True)
    text(s, l + Inches(0.3), y3, w - Inches(0.6), Inches(0.85),
         [[("LOOP principal (Analisar):  comparacao → revisao → comparacao  "
            "— classificar e revisar até consolidar.", 14, WHITE, True, FONT, False)]],
         anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, l, y3 + Inches(1.05), w, Inches(1.4), [
        "Endogeneidade: vencer EXIGE comparar frações — a habilidade-alvo é a mecânica.",
        "Validação semântica: mecânica 'comparacao' (Analisar) coerente com o objetivo (Analisar).",
        "Compila para HTML5 autocontido; reparametrizável p/ 'comparar ângulos', etc.",
    ], size=14, gap=9)


# --------------------------------------------------------------------------- #
# SLIDE 26 — Metodologia DSR
# --------------------------------------------------------------------------- #
def slide_dsr():
    s = new_slide("Metodologia — Design Science Research", "Ciclos de relevância, projeto e rigor")
    l, t, w, h = body_area()
    cols = [
        ("Relevância", "Ambiente", ["Educadores não-programadores",
                                     "Necessidade de jogos alinhados",
                                     "Requisitos do domínio"], SLATE),
        ("Projeto", "Construção & avaliação", ["DSL + parser + validação",
                                               "Pipeline RAG+LLM",
                                               "Compilador + instrumento"], INDIGO),
        ("Rigor", "Base de conhecimento", ["Teoria de DSLs",
                                            "Taxonomia de Bloom",
                                            "Design endógeno · RAG"], RGBColor(0x0D,0x94,0x88)),
    ]
    cw = Inches(3.85); gap = Inches(0.34)
    for i, (ti, sub, items, col) in enumerate(cols):
        x = l + Emu(int(cw + gap)) * i
        box(s, x, t, cw, Inches(4.4), fill=WHITE, line=col, line_w=Pt(2),
            rounded=True, shadow=True)
        box(s, x, t, cw, Inches(0.95), fill=col)
        text(s, x + Inches(0.25), t + Inches(0.1), cw - Inches(0.5), Inches(0.8), [
            [("CICLO DE " + ti.upper(), 12, WHITE, True, FONT, False)],
            [(sub, 14, WHITE, False, FONT, False)],
        ], anchor=MSO_ANCHOR.MIDDLE, space_after=2)
        bullets(s, x + Inches(0.25), t + Inches(1.2), cw - Inches(0.5), Inches(3.0),
                items, size=14, gap=12, marker_color=col)
        if i < 2:
            cy = Emu(int(t + Inches(2.2)))
            arrow(s, x + cw, cy, x + cw + gap, cy, color=GRAY)
    text(s, l, t + Inches(4.6), w, Inches(0.5),
         [[("Artefato avaliado iterativamente; a prova de conceito (Endo-DSL "
            "implementada) já evidencia viabilidade técnica.", 13, GRAY, False, FONT, True)]])


# --------------------------------------------------------------------------- #
# SLIDE 27 — Protocolo experimental
# --------------------------------------------------------------------------- #
def slide_protocolo():
    content_slide(
        "Protocolo de avaliação experimental", "Automático vs. manual (QP4)",
        [
            "Conjunto de cenários pedagógicos (domínio × tópico × nível de Bloom) "
            "definidos a priori.",
            "Braço A — protótipos gerados pelo pipeline automático (origin=auto).",
            "Braço B — protótipos especificados manualmente por autores (origin=manual).",
            "Avaliadores cegos ao braço aplicam o MESMO instrumento de 7 dimensões a "
            "todos os protótipos.",
            "Análise estatística das diferenças por dimensão e do escore agregado "
            "ponderado; teste da hipótese H4.",
            "Reuso/curadoria registrados para avaliar o efeito da biblioteca RAG.",
        ],
        size=17,
        side=[
            "Cenários controlados",
            "Braço A: auto",
            "Braço B: manual",
            "Avaliadores cegos",
            "Mesmo instrumento 7-D",
            "Teste de H4",
        ],
        side_title="Desenho",
    )


# --------------------------------------------------------------------------- #
# SLIDE 28 — Métricas e instrumentação
# --------------------------------------------------------------------------- #
def slide_metricas():
    s = new_slide("Métricas e instrumentação", "O que medimos e como")
    l, t, w, h = body_area()
    groups = [
        ("Geração (pipeline)", RGBColor(0x4F,0x46,0xE5),
         ["Validade na 1ª tentativa", "Nº médio de tentativas",
          "Taxa de rejeição", "Erros mais frequentes"]),
        ("Qualidade pedagógica", RGBColor(0xEA,0x58,0x0C),
         ["Escore por dimensão (1–5)", "Escore agregado ponderado",
          "Concordância entre avaliadores", "Auto vs. manual"]),
        ("Artefato/compilação", RGBColor(0x16,0xA3,0x4A),
         ["Cobertura de mecânicas", "Rastreabilidade objetivo→mecânica",
          "Reparametrizações bem-sucedidas", "Tamanho do HTML5 gerado"]),
    ]
    cw = Inches(3.85); gap = Inches(0.34)
    for i, (ti, col, items) in enumerate(groups):
        x = l + Emu(int(cw + gap)) * i
        box(s, x, t, cw, Inches(4.4), fill=CARD, rounded=True, shadow=True)
        box(s, x, t, cw, Inches(0.7), fill=col)
        text(s, x, t, cw, Inches(0.7), [[(ti, 15, WHITE, True, FONT, False)]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        bullets(s, x + Inches(0.3), t + Inches(0.95), cw - Inches(0.6), Inches(3.3),
                items, size=14, gap=14, marker_color=col)
    text(s, l, t + Inches(4.6), w, Inches(0.5),
         [[("Instrumentação nativa: generation_attempts e evaluations no SQLite; "
            "exportação CSV (RF25) e relatório comparativo (RF26).", 13, GRAY, False, FONT, True)]])


# --------------------------------------------------------------------------- #
# SLIDE 29 — Resultados preliminares / PoC
# --------------------------------------------------------------------------- #
def slide_resultados():
    s = new_slide("Resultados preliminares / prova de conceito", "O que já funciona")
    l, t, w, h = body_area()
    done = [
        "Plataforma implementada em Python 3.11 (stdlib), CLI + Web + API HTTP.",
        "Gramática EBNF + parser recursivo descendente + validação sint./sem. operacionais.",
        "Biblioteca SQLite com componentes canônicos e curadoria funcional.",
        "Pipeline multi-agente com refinamento iterativo e métricas instrumentadas.",
        "Compilador DSL→HTML5 autocontido + reparametrização de domínio.",
        "Estudo de caso 'Comparando Frações' compila e executa no navegador.",
    ]
    box(s, l, t, Inches(7.3), Inches(5.0), fill=NEARWHITE, line=LIGHTGRAY, line_w=Pt(1), rounded=True)
    text(s, l + Inches(0.3), t + Inches(0.15), Inches(6.9), Inches(0.4),
         [[("✓  Já implementado (prova de conceito)", 16, RGBColor(0x16,0xA3,0x4A),
            True, FONT, False)]])
    bullets(s, l + Inches(0.3), t + Inches(0.7), Inches(6.8), Inches(4.2),
            done, size=14.5, gap=11, marker="✓", marker_color=RGBColor(0x16,0xA3,0x4A))
    # painel de indicadores ilustrativos
    x2 = l + Inches(7.6)
    pw = w - Inches(7.6)
    box(s, x2, t, pw, Inches(5.0), fill=SLATE, rounded=True, shadow=True)
    text(s, x2 + Inches(0.3), t + Inches(0.2), pw - Inches(0.6), Inches(0.5),
         [[("Indicadores (ilustrativos)", 14, INDIGO_LT, True, FONT, False)]])
    stats = [("5", "módulos integrados"), ("6", "níveis de Bloom"),
             ("7", "dimensões de avaliação"), ("0", "dependências em runtime")]
    yy = t + Inches(0.9)
    for num, lab in stats:
        text(s, x2 + Inches(0.3), yy, pw - Inches(0.6), Inches(0.95), [
            [(num, 34, WHITE, True, FONT, False)],
            [(lab, 12, RGBColor(0xCB,0xD5,0xE1), False, FONT, False)],
        ], space_after=0)
        yy += Inches(1.0)


# --------------------------------------------------------------------------- #
# SLIDE 30 — Contribuições esperadas
# --------------------------------------------------------------------------- #
def slide_contribuicoes():
    s = new_slide("Contribuições esperadas", "Científicas e técnicas")
    l, t, w, h = body_area()
    cw = Inches(5.95)
    box(s, l, t, cw, Inches(5.0), fill=CARD, rounded=True, shadow=True)
    box(s, l, t, cw, Inches(0.65), fill=INDIGO)
    text(s, l, t, cw, Inches(0.65), [[("Científicas", 16, WHITE, True, FONT, False)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, l + Inches(0.35), t + Inches(0.9), cw - Inches(0.7), Inches(4.0), [
        "DSL com Bloom como construto de 1ª classe e sua semântica de afinidade.",
        "Método de geração de jogos educativos ancorado por DSL (anti-alucinação).",
        "Instrumento de 7 dimensões para comparar protótipos auto vs. manual.",
        "Evidências sobre viabilidade de autoria pedagógica por não-programadores.",
    ], size=15, gap=14)
    x2 = l + cw + Inches(0.33)
    box(s, x2, t, cw, Inches(5.0), fill=CARD, rounded=True, shadow=True)
    box(s, x2, t, cw, Inches(0.65), fill=RGBColor(0x16,0xA3,0x4A))
    text(s, x2, t, cw, Inches(0.65), [[("Técnicas", 16, WHITE, True, FONT, False)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, x2 + Inches(0.35), t + Inches(0.9), cw - Inches(0.7), Inches(4.0), [
        "Plataforma Endo-DSL aberta (CLI + Web + API), sem dependências obrigatórias.",
        "Compilador DSL→HTML5 autocontido com reparametrização de domínio.",
        "Biblioteca curada de componentes pedagógicos reutilizáveis (RAG).",
        "Pipeline multi-agente instrumentado e reprodutível.",
    ], size=15, gap=14, marker_color=RGBColor(0x16,0xA3,0x4A))


# --------------------------------------------------------------------------- #
# SLIDE 31 — Ameaças à validade
# --------------------------------------------------------------------------- #
def slide_ameacas():
    s = new_slide("Ameaças à validade", "E estratégias de mitigação")
    l, t, w, h = body_area()
    rows = [
        ("Construto", "Instrumento de 7-D pode não capturar toda 'qualidade pedagógica'.",
         "Ancorar dimensões em MEEGA+; validação por especialistas."),
        ("Interna", "Geração depende do LLM/seed; variabilidade entre execuções.",
         "Backend reprodutível; múltiplas execuções; métricas agregadas."),
        ("Externa", "Resultados podem não generalizar p/ outros domínios/idades.",
         "Múltiplos cenários (domínio×tópico×Bloom); reparametrização."),
        ("Conclusão", "Tamanho amostral de avaliadores e protótipos.",
         "Poder estatístico planejado; avaliadores cegos ao braço."),
    ]
    headers = ["Tipo", "Ameaça", "Mitigação"]
    cwid = [Inches(2.0), Inches(5.1), Inches(5.1)]
    rh = Inches(0.95); x0 = l; y0 = t + Inches(0.1)
    cx = x0
    for j, hd in enumerate(headers):
        box(s, cx, y0, cwid[j], Inches(0.55), fill=SLATE)
        text(s, cx + Inches(0.15), y0, cwid[j] - Inches(0.3), Inches(0.55),
             [[(hd, 14, WHITE, True, FONT, False)]], anchor=MSO_ANCHOR.MIDDLE)
        cx += cwid[j]
    for i, row in enumerate(rows):
        y = y0 + Inches(0.55) + Emu(int(rh) * i)
        cx = x0
        for j, cell in enumerate(row):
            fill = WHITE if i % 2 == 0 else CARD
            box(s, cx, y, cwid[j], rh, fill=fill, line=LIGHTGRAY, line_w=Pt(0.5))
            tcol = INDIGO if j == 0 else INK
            bold = (j == 0)
            text(s, cx + Inches(0.15), y, cwid[j] - Inches(0.3), rh,
                 [[(cell, 13, tcol, bold, FONT, False)]], anchor=MSO_ANCHOR.MIDDLE)
            cx += cwid[j]


# --------------------------------------------------------------------------- #
# SLIDE 32 — Cronograma (Gantt)
# --------------------------------------------------------------------------- #
def slide_cronograma():
    s = new_slide("Cronograma", "Gantt — 6 semestres (2026–2028)")
    l, t, w, h = body_area()
    tasks = [
        ("Revisão sistemática da literatura", 0, 2),
        ("Projeto da DSL e validação (M1)", 0, 2),
        ("Biblioteca + RAG (M2)", 1, 3),
        ("Pipeline multi-agente (M3)", 2, 4),
        ("Compilador HTML5 (M4)", 2, 4),
        ("Instrumento de avaliação (M5)", 3, 4),
        ("Estudo experimental auto×manual", 4, 5),
        ("Análise de resultados", 4, 6),
        ("Escrita e defesa da tese", 4, 6),
    ]
    nsem = 6
    grid_l = l + Inches(4.5); grid_w = w - Inches(4.6)
    sem_w = Emu(int(grid_w / nsem))
    top = t + Inches(0.55); rh = Inches(0.48)
    # cabeçalho semestres
    for k in range(nsem):
        gx = grid_l + Emu(int(sem_w) * k)
        box(s, gx, top - Inches(0.45), sem_w, Inches(0.4), fill=SLATE_LT,
            line=WHITE, line_w=Pt(0.5))
        text(s, gx, top - Inches(0.45), sem_w, Inches(0.4),
             [[(f"S{k+1}", 12, WHITE, True, FONT, False)]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for i, (name, a, b) in enumerate(tasks):
        y = top + Emu(int(rh) * i)
        text(s, l, y, Inches(4.4), rh, [[(name, 12.5, INK, False, FONT, False)]],
             anchor=MSO_ANCHOR.MIDDLE)
        # grid de fundo
        for k in range(nsem):
            gx = grid_l + Emu(int(sem_w) * k)
            box(s, gx, y + Inches(0.05), sem_w, rh - Inches(0.12),
                fill=NEARWHITE, line=LIGHTGRAY, line_w=Pt(0.5))
        bx = grid_l + Emu(int(sem_w) * a)
        bw = Emu(int(sem_w) * (b - a))
        col = BLOOM[i % len(BLOOM)][1]
        box(s, bx + Inches(0.03), y + Inches(0.08), bw - Inches(0.06), rh - Inches(0.18),
            fill=INDIGO if i < 6 else RGBColor(0xEA,0x58,0x0C), rounded=True)
    text(s, l, top + Emu(int(rh) * len(tasks)) + Inches(0.1), w, Inches(0.4),
         [[("S = semestre · marcos: qualificação (S2) · estudo experimental (S5) · "
            "defesa (S6).", 12, GRAY, False, FONT, True)]])


# --------------------------------------------------------------------------- #
# SLIDE 33 — Plano de escrita da tese
# --------------------------------------------------------------------------- #
def slide_tese():
    s = new_slide("Plano de escrita da tese", "Capítulos")
    l, t, w, h = body_area()
    caps = [
        ("1", "Introdução", "contexto, problema, QP, hipóteses, objetivos"),
        ("2", "Fundamentação", "DSLs, Bloom, design endógeno, RAG+LLM"),
        ("3", "Trabalhos relacionados", "estado da arte e lacuna"),
        ("4", "A linguagem Endo-DSL", "gramática, parser, validação semântica"),
        ("5", "Geração e compilação", "biblioteca/RAG, pipeline, compilador HTML5"),
        ("6", "Avaliação", "instrumento 7-D, protocolo, resultados"),
        ("7", "Conclusão", "contribuições, limitações, trabalhos futuros"),
    ]
    cw = Inches(5.95); ch = Inches(0.95)
    for i, (num, ti, sub) in enumerate(caps):
        col = i % 2
        x = l + col * (cw + Inches(0.33))
        y = t + (i // 2) * (ch + Inches(0.18))
        box(s, x, y, cw, ch, fill=CARD, rounded=True, shadow=True)
        chip(s, x + Inches(0.2), y + Inches(0.23), Inches(0.55), Inches(0.5),
             num, SLATE, size=18)
        text(s, x + Inches(0.95), y + Inches(0.1), cw - Inches(1.1), ch - Inches(0.2), [
            [(ti, 15, INK, True, FONT, False)],
            [(sub, 11.5, GRAY, False, FONT, False)],
        ], anchor=MSO_ANCHOR.MIDDLE, space_after=1)


# --------------------------------------------------------------------------- #
# SLIDE 34 — Conclusão e próximos passos
# --------------------------------------------------------------------------- #
def slide_conclusao():
    content_slide(
        "Conclusão e próximos passos", "Síntese",
        [
            "A Endo-DSL integra DSL pedagógica verificável, geração RAG+LLM ancorada e "
            "avaliação comparável num único ciclo rastreável.",
            "A prova de conceito demonstra viabilidade técnica de ponta a ponta "
            "(do contexto ao protótipo HTML5 executável).",
            "Próximo passo imediato: consolidar a biblioteca de componentes e o pipeline "
            "multi-agente para o estudo experimental.",
            "A seguir: conduzir a avaliação auto×manual com avaliadores cegos e testar "
            "as hipóteses (sobretudo H4).",
            "Finalmente: análise estatística, refinamento do instrumento e redação da tese.",
        ],
        size=18,
        side=[
            "Consolidar biblioteca/RAG",
            "Maturar pipeline multi-agente",
            "Estudo experimental",
            "Testar hipóteses (H4)",
            "Redação e defesa",
        ],
        side_title="Roadmap",
    )


# --------------------------------------------------------------------------- #
# SLIDE 35 — Referências
# --------------------------------------------------------------------------- #
def slide_referencias():
    s = new_slide("Referências", "Seleção (a expandir na tese)")
    l, t, w, h = body_area()
    refs = [
        "Anderson, L. W. & Krathwohl, D. R. (2001). A Taxonomy for Learning, Teaching, "
        "and Assessing: A Revision of Bloom's Taxonomy.",
        "Fowler, M. (2010). Domain-Specific Languages. Addison-Wesley.",
        "Mernik, M., Heering, J. & Sloane, A. M. (2005). When and how to develop "
        "domain-specific languages. ACM Computing Surveys.",
        "Habgood, M. P. J. & Ainsworth, S. E. (2011). Motivating children to learn "
        "effectively: exploring the value of intrinsic integration. JLS.",
        "Lewis, P. et al. (2020). Retrieval-Augmented Generation for Knowledge-"
        "Intensive NLP Tasks. NeurIPS.",
        "Petri, G. & von Wangenheim, C. G. (2017). MEEGA+: an evolution of a model "
        "for the evaluation of educational games. INCoD/UFSC.",
        "Hevner, A. R. et al. (2004). Design Science in Information Systems Research. "
        "MIS Quarterly.",
        "Plass, J. L., Homer, B. D. & Kinzer, C. K. (2015). Foundations of Game-Based "
        "Learning. Educational Psychologist.",
    ]
    tb = s.shapes.add_textbox(l, t, w, Inches(5.2))
    tf = tb.text_frame; tf.word_wrap = True; first = True
    for r in refs:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.space_after = Pt(10)
        run = p.add_run(); run.text = "▸  " + r
        run.font.size = Pt(14.5); run.font.color.rgb = INK; run.font.name = FONT
    text(s, l, t + Inches(5.0), w, Inches(0.4),
         [[("Lista ilustrativa; a revisão sistemática completa integra o Capítulo 3 da "
            "tese.", 12, GRAY, False, FONT, True)]])


# --------------------------------------------------------------------------- #
# SLIDE 36 — Agradecimento / perguntas
# --------------------------------------------------------------------------- #
def slide_obrigado():
    s = prs.slides.add_slide(BLANK)
    _slide_no["n"] += 1
    box(s, 0, 0, SW, SH, fill=SLATE)
    box(s, 0, Inches(4.3), SW, Inches(0.08), fill=INDIGO)
    text(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.6), [
        [("Obrigado.", 56, WHITE, True, FONT, False)],
        [("Perguntas e discussão", 24, INDIGO_LT, False, FONT, False)],
    ], space_after=10)
    text(s, Inches(0.9), Inches(4.6), Inches(11.5), Inches(1.6), [
        [("Caio Azeredo", 18, WHITE, True, FONT, False)],
        [("caiosazeredo@cos.ufrj.br", 16, RGBColor(0x94,0xA3,0xB8), False, FONT, False)],
        [("PESC — COPPE/UFRJ · Exame de Qualificação · 2026",
          14, RGBColor(0x94,0xA3,0xB8), False, FONT, False)],
    ], space_after=6)
    text(s, Inches(0.9), Inches(6.9), Inches(11.5), Inches(0.4),
         [[(FOOTER, 11, INDIGO_LT, False, FONT, False)]])


# --------------------------------------------------------------------------- #
# Montagem
# --------------------------------------------------------------------------- #
def build():
    slide_capa()             # 1
    slide_agenda()           # 2
    slide_contexto()         # 3
    slide_problema()         # 4
    slide_questoes()         # 5
    slide_hipoteses()        # 6
    slide_objetivos()        # 7
    slide_dsls()             # 8
    slide_bloom()            # 9
    slide_endogeno()         # 10
    slide_rag()              # 11
    slide_relacionados()     # 12
    slide_lacuna()           # 13
    slide_visao()            # 14
    slide_arquitetura()      # 15
    slide_mod1()             # 16
    slide_bloom_firstclass() # 17
    slide_validacao()        # 18
    slide_mod2()             # 19
    slide_mod3()             # 20
    slide_mod4()             # 21
    slide_mod5()             # 22
    slide_jornada()          # 23
    slide_interface()        # 24
    slide_estudo()           # 25
    slide_dsr()              # 26
    slide_protocolo()        # 27
    slide_metricas()         # 28
    slide_resultados()       # 29
    slide_contribuicoes()    # 30
    slide_ameacas()          # 31
    slide_cronograma()       # 32
    slide_tese()             # 33
    slide_conclusao()        # 34
    slide_referencias()      # 35
    slide_obrigado()         # 36

    out = Path(__file__).parent / "Endo-DSL-Qualificacao.pptx"
    prs.save(str(out))
    print(f"OK: {out} gerado com {len(prs.slides._sldIdLst)} slides.")


if __name__ == "__main__":
    build()
